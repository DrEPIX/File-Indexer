"""Document text and metadata.

Text extraction feeds the full-text index, so the goal is *searchable* text,
not faithful reproduction: layout, styling and images are discarded and only
the words are kept.

Every backend is optional. Without PyMuPDF a PDF still becomes an asset with
its page count unknown; it simply is not full-text searchable until the extra
is installed. That is the same degradation contract as the rest of the engine —
missing optional dependency means reduced capability, never a failed scan.

**A PDF with no text layer is flagged ``ocr_eligible``, never OCR'd here.**
OCR costs three orders of magnitude more than reading an existing text layer
and needs a model; that makes it a plugin's job. Flagging it is what lets an
OCR plugin find its work later with one indexed query.
"""

from __future__ import annotations

import logging
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from ...util import parse_datetime, to_iso
from .base import Extracted

__all__ = ["extract_document", "extract_text_file", "MAX_SNIFF_BYTES"]

_LOG = logging.getLogger(__name__)

#: A PDF page yielding fewer characters than this is treated as having no text
#: layer. Scanned pages routinely carry a handful of stray glyphs from a
#: header stamp, so zero is the wrong threshold.
_MIN_CHARS_PER_PAGE = 16

#: How much of an unknown file to read when deciding its encoding.
MAX_SNIFF_BYTES = 64 * 1024

_ENCODINGS = ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1")


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    """Clip to ``limit`` characters, reporting whether anything was lost."""
    if limit <= 0 or len(text) <= limit:
        return text, False
    return text[:limit], True


def _clean(text: str) -> str:
    """Collapse the whitespace that layout-driven extraction leaves behind."""
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def extract_text_file(path: Path | str, *, max_chars: int = 2_000_000) -> Extracted:
    """Read a plain-text file, guessing its encoding.

    Tries the encodings that actually occur in the wild, in order of
    likelihood, and finishes with ``latin-1`` — which cannot fail — so a file
    with a broken encoding still contributes searchable words rather than being
    dropped.
    """
    result = Extracted(extractor="text")
    target = Path(path)
    raw = b""
    try:
        with target.open("rb") as handle:
            raw = handle.read(max(max_chars * 4, MAX_SNIFF_BYTES))
    except OSError as exc:
        result.warnings.append(f"unreadable: {exc}")
        return result

    text = ""
    for encoding in _ENCODINGS:
        try:
            text = raw.decode(encoding)
            result.raw.setdefault("text", {})["encoding"] = encoding
            break
        except (UnicodeDecodeError, LookupError):
            continue

    text, truncated = _truncate(_clean(text), max_chars)
    result.text = text
    result.text_truncated = truncated
    result.payload["word_count"] = len(text.split())
    return result


def _extract_pdf(path: Path, max_chars: int) -> Extracted:
    """PDF via PyMuPDF, flagging scans for a downstream OCR plugin."""
    result = Extracted(extractor="pymupdf")
    try:
        import pymupdf  # type: ignore[import-not-found]
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore[import-not-found, no-redef]
        except ImportError:
            result.warnings.append("PyMuPDF not installed; PDF text not extracted")
            result.ocr_eligible = False
            return result

    try:
        with pymupdf.open(str(path)) as document:
            result.payload["page_count"] = document.page_count
            first = document.load_page(0) if document.page_count else None
            if first is not None:
                rect = first.rect
                result.payload["width"] = int(rect.width)
                result.payload["height"] = int(rect.height)

            chunks: list[str] = []
            budget = max_chars
            for page in document:
                if budget <= 0:
                    break
                page_text = page.get_text("text") or ""
                chunks.append(page_text)
                budget -= len(page_text)

            metadata = dict(document.metadata or {})
            result.raw["pdf"] = {k: v for k, v in metadata.items() if v}
            if metadata.get("title"):
                result.raw.setdefault("pdf", {})["title"] = metadata["title"]
            parsed, tz = parse_datetime(_pdf_date(metadata.get("creationDate")))
            if parsed is not None:
                result.captured_at = to_iso(parsed)
                result.captured_at_tz = tz
                result.captured_at_source = "container"

            text = _clean("\n".join(chunks))
            text, truncated = _truncate(text, max_chars)
            result.text = text
            result.text_truncated = truncated
            result.payload["word_count"] = len(text.split())

            pages = result.payload.get("page_count") or 1
            # No text layer worth the name: this is a scan. Flag it and move
            # on — the OCR plugin will find it via ocr_eligible_asset_ids().
            result.ocr_eligible = len(text) < _MIN_CHARS_PER_PAGE * pages
    except Exception as exc:  # noqa: BLE001 - a malformed PDF must not fail a scan
        result.warnings.append(f"PDF extraction failed: {exc}")
        result.ocr_eligible = True
    return result


def _pdf_date(value: Any) -> str | None:
    """``D:20190407142211+01'00'`` -> something :func:`parse_datetime` accepts."""
    if not value:
        return None
    text = str(value).strip()
    if text.startswith("D:"):
        text = text[2:]
    if len(text) >= 14 and text[:14].isdigit():
        stamp = text[:14]
        return (
            f"{stamp[0:4]}:{stamp[4:6]}:{stamp[6:8]} "
            f"{stamp[8:10]}:{stamp[10:12]}:{stamp[12:14]}"
        )
    return text or None


def _extract_docx(path: Path, max_chars: int) -> Extracted:
    result = Extracted(extractor="python-docx")
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError:
        result.warnings.append("python-docx not installed; .docx text not extracted")
        return result
    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        core = document.core_properties
        result.raw["docx"] = {
            "author": core.author,
            "title": core.title,
            "subject": core.subject,
            "keywords": core.keywords,
            "revision": core.revision,
        }
        if core.created is not None:
            result.captured_at = to_iso(core.created)
            result.captured_at_source = "container"
        text, truncated = _truncate(_clean("\n".join(parts)), max_chars)
        result.text = text
        result.text_truncated = truncated
        result.payload["word_count"] = len(text.split())
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(f"docx extraction failed: {exc}")
    return result


def _extract_xlsx(path: Path, max_chars: int) -> Extracted:
    result = Extracted(extractor="openpyxl")
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError:
        result.warnings.append("openpyxl not installed; spreadsheet text not extracted")
        return result
    try:
        # read_only + data_only: values not formulas, and streamed rather than
        # building an object graph for a 100 MB workbook.
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        budget = max_chars
        for sheet in workbook.worksheets:
            parts.append(str(sheet.title))
            for row in sheet.iter_rows(values_only=True):
                if budget <= 0:
                    break
                line = " ".join(str(v) for v in row if v is not None)
                if line:
                    parts.append(line)
                    budget -= len(line)
            if budget <= 0:
                break
        result.payload["page_count"] = len(workbook.worksheets)
        result.raw["xlsx"] = {"sheets": [s.title for s in workbook.worksheets]}
        workbook.close()
        text, truncated = _truncate(_clean("\n".join(parts)), max_chars)
        result.text = text
        result.text_truncated = truncated
        result.payload["word_count"] = len(text.split())
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(f"xlsx extraction failed: {exc}")
    return result


def _extract_pptx(path: Path, max_chars: int) -> Extracted:
    """PowerPoint without a dependency on python-pptx.

    A .pptx is a ZIP of XML slides; pulling every ``<a:t>`` text run out of
    them gives exactly the words, which is all the index wants. Adding another
    optional dependency to do the same job is not worth it.
    """
    result = Extracted(extractor="pptx-xml")
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
    try:
        with zipfile.ZipFile(path) as archive:
            slides = sorted(
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
            parts: list[str] = []
            for name in slides:
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError:
                    continue
                parts.extend(node.text for node in root.iter(namespace) if node.text)
            result.payload["page_count"] = len(slides)
            text, truncated = _truncate(_clean("\n".join(parts)), max_chars)
            result.text = text
            result.text_truncated = truncated
            result.payload["word_count"] = len(text.split())
    except (zipfile.BadZipFile, OSError) as exc:
        result.warnings.append(f"pptx extraction failed: {exc}")
    return result


def _extract_epub(path: Path, max_chars: int) -> Extracted:
    result = Extracted(extractor="ebooklib")
    try:
        import ebooklib  # type: ignore[import-not-found]
        from ebooklib import epub  # type: ignore[import-not-found]
    except ImportError:
        result.warnings.append("EbookLib not installed; EPUB text not extracted")
        return result
    try:
        import re

        book = epub.read_epub(str(path))
        parts: list[str] = []
        budget = max_chars
        chapters = 0
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            chapters += 1
            if budget <= 0:
                continue
            html = item.get_content().decode("utf-8", errors="replace")
            stripped = re.sub(r"<[^>]+>", " ", html)
            parts.append(stripped)
            budget -= len(stripped)
        result.payload["page_count"] = chapters
        titles = book.get_metadata("DC", "title")
        creators = book.get_metadata("DC", "creator")
        result.raw["epub"] = {
            "title": titles[0][0] if titles else None,
            "creator": creators[0][0] if creators else None,
        }
        text, truncated = _truncate(_clean(" ".join(parts)), max_chars)
        result.text = text
        result.text_truncated = truncated
        result.payload["word_count"] = len(text.split())
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(f"EPUB extraction failed: {exc}")
    return result


def _extract_html(path: Path, max_chars: int) -> Extracted:
    """Strip tags from HTML. Same reasoning as pptx: no new dependency."""
    import re

    result = extract_text_file(path, max_chars=max_chars * 4)
    if result.text:
        without_script = re.sub(
            r"<(script|style)[^>]*>.*?</\1>", " ", result.text, flags=re.S | re.I
        )
        text, truncated = _truncate(_clean(re.sub(r"<[^>]+>", " ", without_script)), max_chars)
        result.text = text
        result.text_truncated = truncated or result.text_truncated
        result.payload["word_count"] = len(text.split())
    result.extractor = "html"
    return result


#: mime -> backend. Dispatch is on the *sniffed* mime, never the extension.
_BACKENDS = {
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _extract_pptx,
    "application/epub+zip": _extract_epub,
    "text/html": _extract_html,
}


def extract_document(
    path: Path | str, *, mime_type: str | None = None, max_chars: int = 2_000_000
) -> Extracted:
    """Extract text and metadata from a document, dispatching on mime type."""
    target = Path(path)
    backend = _BACKENDS.get(mime_type or "")
    if backend is not None:
        return backend(target, max_chars)
    if (mime_type or "").startswith("text/") or mime_type in (
        "application/json",
        "application/x-subrip",
    ):
        return extract_text_file(target, max_chars=max_chars)

    result = Extracted(extractor="none")
    result.warnings.append(f"no text extractor for {mime_type or 'unknown type'}")
    return result
